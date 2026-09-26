"""Build the interim paper from docs/paper/paper.md into HTML, PDF (headless Edge/Chrome) and DOCX.

Numbers in tables come straight from results/*.csv and the scenario configuration, via
{{placeholders}} in the Markdown source, so the paper can never drift from the experiments.

    python tools/build_paper.py            # writes docs/paper/paper.{html,pdf,docx}
"""

from __future__ import annotations

import html
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import numpy as np  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt, RGBColor, Inches  # noqa: E402

from satsched.experiments import read_csv  # noqa: E402
from satsched.geometry import load_geometry  # noqa: E402
from satsched.scenarios import NOMINAL, SCENARIOS  # noqa: E402

PAPER = ROOT / "docs" / "paper"
RESULTS = ROOT / "results"
POLICY = {"value_aware": "Value-aware (ours)", "bandwidth_rules": "Bandwidth-aware rules",
          "priority_rules": "Priority rules", "process_all": "Process-all", "bent_pipe": "Bent-pipe"}
VARIANT = {"full_engine": "Full engine", "no_prices": "No shadow prices (all λ = 0)",
           "no_storage_price": "No storage price", "no_queue_pass": "No queue-aware re-valuation",
           "no_battery_guard": "No battery guard"}
SCEN = {"nominal": "Nominal", "energy_starved": "Energy-starved", "downlink_starved": "Downlink-starved",
        "storage_tight": "Storage-tight", "compute_starved": "Compute-starved", "event_surge": "Event surge"}

# ---------------------------------------------------------------- placeholder tables


def _pm(row: dict, key: str, digits: int = 1, compact: bool = False) -> str:
    sep = "±" if compact else " ± "
    return f"{float(row[key]):.{digits}f}{sep}{float(row[key + '_ci95']):.{digits}f}"


def _table(header: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(lines)


def results_table() -> str:
    summary = read_csv(RESULTS / "summary.csv")
    rows = []
    for s in SCEN:
        for p in POLICY:
            r = next((x for x in summary if x["scenario"] == s and x["policy"] == p), None)
            if r is None:
                continue
            name = f"**{POLICY[p]}**" if p == "value_aware" else POLICY[p]
            rows.append([SCEN[s] if p == "value_aware" else "", name] + [
                _pm(r, key, d, compact=True) for key, d in (
                    ("value_score_pct", 1), ("completion_rate_pct", 1), ("energy_total_wh", 0), ("min_soc_pct", 0),
                    ("latency_mean_min", 0), ("alert_latency_mean_min", 0), ("downlink_util_pct", 0),
                    ("storage_peak_pct", 0))])
    return _table(["Scenario", "Strategy", "Value %", "Completed %", "Energy Wh", "Min battery %", "Latency min",
                   "Alert lat. min", "Downlink %", "Peak storage %"], rows)


def ablation_table() -> str:
    path = RESULTS / "ablation_summary.csv"
    if not path.exists():
        return "*(ablation results not available — run `python -m satsched.cli ablation`)*"
    summary = read_csv(path)
    rows = []
    for v in VARIANT:
        cells = [VARIANT[v]]
        for s in SCEN:
            r = next((x for x in summary if x["scenario"] == s and x["policy"] == v), None)
            cells.append(_pm(r, "value_score_pct", compact=True) if r else "–")
        r_e = next((x for x in summary if x["scenario"] == "energy_starved" and x["policy"] == v), None)
        cells.append(f"{float(r_e['time_below_reserve_pct']):.1f}" if r_e else "–")
        rows.append(cells)
    return _table(["Variant"] + [SCEN[s] for s in SCEN] + ["Below reserve % (energy-starved)"], rows)


def types_table() -> str:
    rows = []
    for k in NOMINAL.data_types:
        proc = "–" if not k.processable else f"{k.processor.value.upper()} {k.proc_time_s:.0f} s"
        rows.append([k.name.replace("_", " "), f"{k.raw_size_mb:.0f}", f"{k.base_value:g}", f"{k.rate_per_orbit:g}",
                     f"{k.half_life_s / 3600:g} h / {k.deadline_s / 3600:g} h", f"{k.p_useful:g}", proc,
                     "–" if not k.processable else f"{k.product_ratio:g}", "–" if not k.processable else f"{k.retention:g}",
                     f"{k.ground_delay_s / 3600:g} h"])
    return _table(["Data type", "Raw MB", "Base value", "Per orbit", "Half-life / deadline", "P(useful)",
                   "Processing", "Product ratio", "Retention", "Ground delay"], rows)


def _passes_per_day(orbit) -> str:
    if not orbit.geometry:
        return f"{orbit.clusters_per_day * orbit.passes_per_cluster}"
    geo = load_geometry(orbit.geometry)
    n = sum(1 for m, *_ in geo.passes if m == orbit.min_elevation_deg)
    return f"{n / geo.meta['days']:.1f} (≥ {orbit.min_elevation_deg:g}°)"


def scenarios_table() -> str:
    rows = []
    for name, sc in SCENARIOS.items():
        sat, orb = sc.satellite, sc.orbit
        rows.append([SCEN[name], f"{sat.solar_w:g} W", f"{sat.battery_wh:g} Wh", f"{sat.storage_mb / 1000:g} GB",
                     _passes_per_day(orb), f"{orb.downlink_mb_s:g} MB/s", sc.description])
    return _table(["Scenario", "Solar", "Battery", "Storage", "Passes/day", "Downlink", "What binds"], rows)


def synthetic_table() -> str:
    """Robustness check: value captured with the earlier synthetic orbit geometry (same seeds, same engine)."""
    path = RESULTS / "synthetic" / "summary.csv"
    if not path.exists():
        return "*(synthetic-geometry results not available)*"
    summary = [{**r, "scenario": str(r["scenario"]).removesuffix("_synthetic")} for r in read_csv(path)]
    rows = []
    for p in POLICY:
        cells = [POLICY[p]]
        for s in SCEN:
            r = next((x for x in summary if x["scenario"] == s and x["policy"] == p), None)
            cells.append(_pm(r, "value_score_pct", compact=True) if r else "–")
        rows.append(cells)
    return _table(["Strategy"] + [SCEN[s] for s in SCEN], rows)


def paired_table() -> str:
    """Per-seed paired differences in value captured (value-aware minus baseline)."""
    runs = read_csv(RESULTS / "runs.csv")
    by = {(r["scenario"], r["policy"], int(r["seed"])): float(r["value_score_pct"]) for r in runs}
    seeds = sorted({int(r["seed"]) for r in runs})
    rows = []
    for s in SCEN:
        cells = [SCEN[s]]
        for ref in ("bandwidth_rules", "priority_rules", "process_all"):
            diffs = np.array([by[(s, "value_aware", k)] - by[(s, ref, k)] for k in seeds if (s, ref, k) in by])
            ci = 1.96 * diffs.std(ddof=1) / np.sqrt(len(diffs))
            cells += [f"{diffs.mean():+.2f}±{ci:.2f}", f"{int((diffs > 0).sum())}/{len(diffs)}"]
        rows.append(cells)
    return _table(["Scenario", "Δ vs bandwidth rules", "Wins", "Δ vs priority rules", "Wins", "Δ vs process-all",
                   "Wins"], rows)


BINDING = [  # (stress scenario, resource, metric, label)
    ("energy_starved", "Energy", "min_soc_pct", "lowest battery (%)"),
    ("downlink_starved", "Downlink", "downlink_util_pct", "pass capacity used (%)"),
    ("storage_tight", "Storage", "storage_peak_pct", "peak storage (%)"),
    ("compute_starved", "Compute", "gpu_util_pct", "GPU busy (%)"),
]


def binding_table() -> str:
    """Scenario validation: each stress scenario's resource, nominal -> stressed, per strategy."""
    summary = read_csv(RESULTS / "summary.csv")
    by = {(r["scenario"], r["policy"]): r for r in summary}
    pols = ("process_all", "priority_rules", "bandwidth_rules", "value_aware")
    rows = []
    for scen, resource, metric, label in BINDING:
        cells = [SCEN[scen], f"{resource}: {label}"]
        for p in pols:
            cells.append(f"{float(by[('nominal', p)][metric]):.0f} → **{float(by[(scen, p)][metric]):.0f}**")
        rows.append(cells)
    return _table(["Scenario", "Binding resource", "Process-all", "Priority rules", "Bandwidth rules", "Value-aware"],
                  rows)


PLACEHOLDERS = {
    "binding_table": binding_table,
    "synthetic_table": synthetic_table,
    "paired_table": paired_table,
    "results_table": results_table,
    "ablation_table": ablation_table,
    "types_table": types_table,
    "scenarios_table": scenarios_table,
}


def fill(text: str) -> str:
    return re.sub(r"\{\{(\w+)\}\}", lambda m: PLACEHOLDERS[m.group(1)](), text)


# ---------------------------------------------------------------- tiny markdown parser

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|<sub>.+?</sub>|<sup>.+?</sup>)")


def blocks(md: str) -> list[tuple[str, object]]:
    out: list[tuple[str, object]] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            out.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            out.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            out.append(("h3", line[4:].strip()))
        elif line.startswith("$$"):
            out.append(("eq", line.strip("$ ").strip()))
        elif m := re.match(r"!\[(.*)\]\((.+)\)", line):
            out.append(("fig", (m.group(1), m.group(2))))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue
        elif re.match(r"^(- |\d+\. )", line):
            ordered = bool(re.match(r"^\d+\. ", line))
            items = []
            while i < len(lines) and re.match(r"^(- |\d+\. )", lines[i]):
                items.append(re.sub(r"^(- |\d+\. )", "", lines[i]).strip())
                i += 1
            out.append(("ol" if ordered else "ul", items))
            continue
        elif line.startswith("> "):
            out.append(("meta", line[2:].strip()))
        else:
            para = [line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\$\$|- |\d+\. |> )", lines[i]):
                para.append(lines[i].strip())
                i += 1
            out.append(("p", " ".join(para)))
            continue
        i += 1
    return out


def inline_html(text: str) -> str:
    parts = []
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            parts.append(f"<strong>{inline_html(tok[2:-2])}</strong>")
        elif tok.startswith("`"):
            parts.append(f"<code>{html.escape(tok[1:-1])}</code>")
        elif tok.startswith("<sub>") or tok.startswith("<sup>"):
            tag = tok[1:4]
            parts.append(f"<{tag}>{inline_html(tok[5:-6])}</{tag}>")
        elif tok.startswith("*") and tok.endswith("*"):
            parts.append(f"<em>{inline_html(tok[1:-1])}</em>")
        else:
            parts.append(html.escape(tok))
    return "".join(parts)


CSS = """
@page { size: A4; margin: 20mm 18mm 20mm 18mm; }
body { font: 10.5pt/1.45 Cambria, Georgia, 'Times New Roman', serif; color: #111; max-width: 180mm; margin: 0 auto; }
h1 { font: 700 20pt/1.2 'Segoe UI', Calibri, Arial, sans-serif; margin: 0 0 6pt; text-align: center; }
.meta { text-align: center; color: #333; margin: 0 0 2pt; font-size: 10pt; }
h2 { font: 700 12.5pt/1.3 'Segoe UI', Calibri, Arial, sans-serif; margin: 16pt 0 5pt; break-after: avoid; }
h3 { font: 600 11pt/1.3 'Segoe UI', Calibri, Arial, sans-serif; margin: 11pt 0 4pt; break-after: avoid; }
p { margin: 0 0 6pt; text-align: justify; hyphens: auto; }
ul, ol { margin: 0 0 6pt 16pt; padding: 0; } li { margin: 0 0 2pt; }
code { font: 9pt Consolas, 'Cascadia Mono', monospace; background: #f3f3f1; padding: 0 2px; }
.eq { display: flex; justify-content: center; align-items: center; margin: 6pt 0 8pt; font-style: italic; font-size: 11pt; position: relative; }
.eq .num { position: absolute; right: 0; font-style: normal; }
figure { margin: 8pt 0 10pt; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; }
figcaption { font-size: 9pt; color: #333; margin-top: 3pt; text-align: left; }
table { border-collapse: collapse; width: 100%; font-size: 8.3pt; margin: 4pt 0 10pt; break-inside: avoid; }
th, td { border-bottom: 0.6pt solid #bbb; padding: 2.5pt 4pt; text-align: left; vertical-align: top; }
th { border-top: 1pt solid #111; border-bottom: 1pt solid #111; font-weight: 700; }
.abstract { border-left: 2.5pt solid #2a78d6; padding: 2pt 0 2pt 10pt; margin: 10pt 0 12pt; }
.tcap { font-size: 9pt; color: #333; margin: 8pt 0 2pt; }
"""


def to_html(md: str) -> str:
    body = []
    fig_n = eq_n = 0
    in_abstract = False
    for kind, val in blocks(md):
        if kind == "title":
            body.append(f"<h1>{inline_html(val)}</h1>")
        elif kind == "meta":
            body.append(f'<p class="meta">{inline_html(val)}</p>')
        elif kind == "h2":
            if in_abstract:
                body.append("</div>")
                in_abstract = False
            if val.lower() == "abstract":
                body.append('<div class="abstract"><h3>Abstract</h3>')
                in_abstract = True
            else:
                body.append(f"<h2>{inline_html(val)}</h2>")
        elif kind == "h3":
            body.append(f"<h3>{inline_html(val)}</h3>")
        elif kind == "p":
            cls = ' class="tcap"' if str(val).startswith("**Table") else ""
            body.append(f"<p{cls}>{inline_html(val)}</p>")
        elif kind in ("ul", "ol"):
            body.append(f"<{kind}>" + "".join(f"<li>{inline_html(x)}</li>" for x in val) + f"</{kind}>")
        elif kind == "eq":
            eq_n += 1
            body.append(f'<div class="eq"><span>{inline_html(val)}</span><span class="num">({eq_n})</span></div>')
        elif kind == "fig":
            fig_n += 1
            cap, src = val
            body.append(f'<figure><img src="{html.escape(src)}" alt=""><figcaption><strong>Figure {fig_n}.</strong> '
                        f"{inline_html(cap)}</figcaption></figure>")
        elif kind == "table":
            head, *rows = val
            body.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in head) + "</tr></thead><tbody>"
                        + "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in rows)
                        + "</tbody></table>")
    if in_abstract:
        body.append("</div>")
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Interim paper</title>"
            f"<style>{CSS}</style></head><body>{''.join(body)}</body></html>")


# ---------------------------------------------------------------- docx


def _runs(par, text: str, size: float | None = None, **style: bool) -> None:
    """Add formatted runs; bold/italic/sub/sup nest (e.g. *p<sub>i</sub>*)."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            _runs(par, tok[2:-2], size, **{**style, "bold": True})
            continue
        if tok.startswith("<sub>"):
            _runs(par, tok[5:-6], size, **{**style, "subscript": True})
            continue
        if tok.startswith("<sup>"):
            _runs(par, tok[5:-6], size, **{**style, "superscript": True})
            continue
        if tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            _runs(par, tok[1:-1], size, **{**style, "italic": True})
            continue
        code = tok.startswith("`")
        r = par.add_run(tok[1:-1] if code else tok)
        if code:
            r.font.name = "Consolas"
        r.bold = style.get("bold") or None
        r.italic = style.get("italic") or None
        r.font.subscript = style.get("subscript") or None
        r.font.superscript = style.get("superscript") or None
        if size:
            r.font.size = Pt(size)


def to_docx(md: str, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)
    fig_n = eq_n = 0
    for kind, val in blocks(md):
        if kind == "title":
            p = doc.add_heading(level=0)
            _runs(p, val)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "meta":
            p = doc.add_paragraph()
            _runs(p, val)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "h2":
            _runs(doc.add_heading(level=1), val)
        elif kind == "h3":
            _runs(doc.add_heading(level=2), val)
        elif kind == "p":
            _runs(doc.add_paragraph(), val)
        elif kind in ("ul", "ol"):
            for item in val:
                _runs(doc.add_paragraph(style="List Number" if kind == "ol" else "List Bullet"), item)
        elif kind == "eq":
            eq_n += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, f"*{val}*    ({eq_n})")
        elif kind == "fig":
            fig_n += 1
            cap, src = val
            img = (PAPER / src).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph()
            _runs(p, f"**Figure {fig_n}.** {cap}", size=9)
        elif kind == "table":
            head, *rows = val
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Light List Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(t.rows[0].cells, head):
                cell.text = ""
                _runs(cell.paragraphs[0], f"**{text}**" if not text.startswith("**") else text, size=8)
            for r in rows:
                cells = t.add_row().cells
                for cell, text in zip(cells, r):
                    cell.text = ""
                    _runs(cell.paragraphs[0], text, size=8)
            doc.add_paragraph()
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.color is not None and r.font.color.rgb is None:
                r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    doc.save(str(path))


def find_browser() -> str | None:
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return shutil.which("msedge") or shutil.which("google-chrome") or shutil.which("chromium")


def main() -> int:
    md = fill((PAPER / "paper.md").read_text(encoding="utf-8"))
    (PAPER / "paper.filled.md").write_text(md, encoding="utf-8")
    html_path = PAPER / "paper.html"
    html_path.write_text(to_html(md), encoding="utf-8")
    to_docx(md, PAPER / "paper.docx")
    browser = find_browser()
    if browser:
        subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={PAPER / 'paper.pdf'}", html_path.as_uri()], check=False,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    else:
        print("no Edge/Chrome found: open paper.html and print to PDF manually")
    print("built:", ", ".join(p.name for p in PAPER.glob("paper.*")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
